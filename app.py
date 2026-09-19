"""
================================================================================
 APPLICATION STREAMLIT - RECOMMANDATIONS PORTEFEUILLE ML MASI20
 (reprend exactement la logique de pipeline_ml_masi20_v6.py, sans feuille Macro)
================================================================================
Lancer en local :  streamlit run app.py
"""

import io
import warnings
warnings.filterwarnings("ignore")

import numpy as np
import pandas as pd
import matplotlib.pyplot as plt
import streamlit as st

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from sklearn.ensemble import RandomForestRegressor
from sklearn.svm import SVR
from sklearn.neural_network import MLPRegressor
from sklearn.preprocessing import StandardScaler
from sklearn.metrics import mean_squared_error, mean_absolute_error, r2_score
from sklearn.base import clone

try:
    from xgboost import XGBRegressor
    from lightgbm import LGBMRegressor
    from catboost import CatBoostRegressor
    PACKAGES_OK = True
except ImportError as e:
    PACKAGES_OK = False
    IMPORT_ERROR = str(e)

st.set_page_config(page_title="Recommandations Portefeuille ML - MASI20", layout="wide")

# ==============================================================================
# PARAMETRES PAR DEFAUT
# ==============================================================================
TARGET_COL = "Rendement M+1 (Cible)"
TOP_N_DEFAULT = 10
FEATURES_TECHNIQUE = [
    "Momentum_3M", "Mom_6M", "Mom_12M",
    "MA3_ratio", "MA6_ratio", "MA12_ratio",
    "RSI_6M", "RSI_12M",
    "EMA_fast3_ratio", "EMA_slow6_ratio",
    "MACD_line_pct", "MACD_signal_pct", "MACD_hist_pct",
    "RollMax12_ratio", "Drawdown", "MaxDD_12M",
]
FEATURES_FOND = ["ROA %", "ROE %", "Marge Nette %", "Marge EBITDA %",
                  "Croissance CA %", "P/E", "P/B", "Croissance RNPG", "Croissance EBITDA"]


def construire_modeles():
    return {
        "RandomForest": RandomForestRegressor(n_estimators=300, max_depth=5, min_samples_leaf=5, random_state=42, n_jobs=-1),
        "XGBoost": XGBRegressor(n_estimators=200, learning_rate=0.02, max_depth=4, subsample=0.8, colsample_bytree=0.8, random_state=42, verbosity=0),
        "LightGBM": LGBMRegressor(n_estimators=200, learning_rate=0.02, num_leaves=15, min_child_samples=5, random_state=42, verbose=-1, importance_type="gain"),
        "CatBoost": CatBoostRegressor(iterations=200, learning_rate=0.05, depth=4, random_seed=42, verbose=0),
        "SVR": SVR(kernel="rbf", C=1, epsilon=0.001, gamma="scale"),
        "MLP": MLPRegressor(hidden_layer_sizes=(64, 32), activation="relu", learning_rate_init=0.001, max_iter=500, random_state=42, early_stopping=True, validation_fraction=0.15),
    }


MODELES_ARBRES = ["RandomForest", "XGBoost", "LightGBM", "CatBoost"]
MODELES_A_SCALER = ["SVR", "MLP"]

# ==============================================================================
# PREPARATION DES DONNEES
# ==============================================================================
@st.cache_data(show_spinner="Preparation des donnees...")
def preparer_donnees(fichier_bytes):
    xls = pd.ExcelFile(io.BytesIO(fichier_bytes))
    df = pd.read_excel(xls, "Technique")
    df_fond = pd.read_excel(xls, "Fondamental")

    for d in (df, df_fond):
        d.columns = d.columns.astype(str).str.replace("\ufeff", "", regex=False).str.strip()

    # Nettoyage Fondamental : retirer les lignes de notes/sources (pas des donnees ticker)
    motif_periode_valide = df_fond["Période"].astype(str).str.match(r"^S[12]\s\d{4}$", na=False)
    df_fond = df_fond[motif_periode_valide].reset_index(drop=True)

    df["Ticker"] = df["Ticker"].ffill()
    tickers_source = sorted(df["Ticker"].unique())

    # Detection automatique du bug de date (jour = mois reel) au lieu de le supposer
    df["Date"] = pd.to_datetime(df["Date"], errors="coerce")
    df = df.dropna(subset=["Date"])
    jours, mois = df["Date"].dt.day, df["Date"].dt.month
    bug_detecte = (mois.nunique() == 1) and jours.between(1, 12).all() and jours.nunique() > 1
    if bug_detecte:
        df["Date"] = pd.to_datetime(dict(year=df["Date"].dt.year, month=df["Date"].dt.day, day=1))
    else:
        df["Date"] = df["Date"].values.astype("datetime64[M]")
    df = df.sort_values(["Ticker", "Date"]).reset_index(drop=True)
    df["Annee"] = df["Date"].dt.year
    df["Mois"] = df["Date"].dt.month

    # Cible
    if TARGET_COL not in df.columns:
        df[TARGET_COL] = df.groupby("Ticker")["Price"].shift(-1) / df["Price"] - 1

    # SECURITE : la derniere observation de chaque ticker ne peut jamais avoir un
    # rendement M+1 reellement connu (pas de prix du mois suivant). Certains fichiers
    # source laissent une valeur non-nulle a cet endroit (ex: 0.0 au lieu de vide) -
    # on la force a NaN quoi qu'il arrive pour ne jamais confondre un mois "a venir"
    # avec un mois "realise".
    idx_dernieres_lignes = df.sort_values(["Ticker", "Date"]).groupby("Ticker").tail(1).index
    df.loc[idx_dernieres_lignes, TARGET_COL] = np.nan

    # Normalisation technique (ratio / prix)
    price_safe = df["Price"].replace(0, np.nan)
    df["MA3_ratio"] = df["Moyenne Mobile_3M"] / price_safe
    df["MA6_ratio"] = df["MM_6M"] / price_safe
    df["MA12_ratio"] = df["MM12M"] / price_safe
    df["EMA_fast3_ratio"] = df["EMA_fast3"] / price_safe
    df["EMA_slow6_ratio"] = df["EMA_slow6"] / price_safe
    df["MACD_line_pct"] = df["MACD_line"] / price_safe
    df["MACD_signal_pct"] = df["MACD_signal3"] / price_safe
    df["MACD_hist_pct"] = df["MACD_hist"] / price_safe
    df["RollMax12_ratio"] = df["RollMax_12"] / price_safe

    # Fondamental (semestriel, decale d'une periode pour eviter le look-ahead)
    def parse_semestre(s):
        sem, annee = s.split()
        return int(annee), sem

    def semestre_suivant(annee, sem, n=1):
        total = annee * 2 + (0 if sem == "S1" else 1) + n
        return total // 2, ("S1" if total % 2 == 0 else "S2")

    df_fond[["Annee_pub", "Sem_pub"]] = df_fond["Période"].apply(lambda s: pd.Series(parse_semestre(s)))
    df_fond[["Annee_use", "Sem_use"]] = df_fond.apply(lambda r: pd.Series(semestre_suivant(r["Annee_pub"], r["Sem_pub"])), axis=1)
    df["Semestre"] = np.where(df["Mois"] <= 6, "S1", "S2")

    # Fusion en REPORT VERS L'AVANT (forward-fill), pas en correspondance exacte :
    # chaque mois utilise le DERNIER rapport fondamental deja publiable a cette
    # date, meme s'il date de plusieurs semestres (ex: S2-2025 reutilise pour
    # tout 2026 tant qu'aucun rapport 2026 n'est ajoute), au lieu de laisser les
    # features vides quand aucune donnee plus recente n'existe.
    def indice_periode(annee, sem):
        return annee * 2 + np.where(np.asarray(sem) == "S1", 0, 1)

    df["Periode_idx"] = indice_periode(df["Annee"], df["Semestre"])
    df_fond["Periode_idx_use"] = indice_periode(df_fond["Annee_use"], df_fond["Sem_use"])

    df_fond_pour_merge = df_fond[["Ticker", "Periode_idx_use"] + FEATURES_FOND].copy()
    df = df.sort_values("Periode_idx")
    df_fond_pour_merge = df_fond_pour_merge.sort_values("Periode_idx_use")
    df = pd.merge_asof(
        df, df_fond_pour_merge,
        left_on="Periode_idx", right_on="Periode_idx_use", by="Ticker", direction="backward",
    )
    df = df.sort_values(["Ticker", "Date"]).reset_index(drop=True)
    df = df.drop(columns=["Periode_idx", "Periode_idx_use"])

    all_features = FEATURES_TECHNIQUE + FEATURES_FOND

    # df_features : lignes dont les FEATURES sont completes, cible connue ou non
    # (necessaire pour predire un mois pas encore realise).
    # df_ml : sous-ensemble dont la cible est EN PLUS connue - seule base utilisee
    # pour ENTRAINER les modeles.
    df_features = df.dropna(subset=all_features).copy()
    df_ml = df_features.dropna(subset=[TARGET_COL]).copy()

    tickers_fond = sorted(df_fond["Ticker"].unique())
    tickers_ml = sorted(df_ml["Ticker"].unique())
    tickers_exclus = sorted(set(tickers_source) - set(tickers_ml))

    return df_features, df_ml, all_features, tickers_source, tickers_fond, tickers_exclus


# ==============================================================================
# SELECTION DU MEILLEUR MODELE (train <= annee_train_max, validation = annee_validation)
# ==============================================================================
@st.cache_data(show_spinner="Comparaison des modeles sur la validation...")
def selectionner_modele(df_ml, all_features, annee_train_max, annee_validation):
    train_mask = df_ml["Annee"] <= annee_train_max
    val_mask = df_ml["Annee"] == annee_validation

    X_train, y_train = df_ml.loc[train_mask, all_features], df_ml.loc[train_mask, TARGET_COL]
    X_val, y_val = df_ml.loc[val_mask, all_features], df_ml.loc[val_mask, TARGET_COL]

    scaler = StandardScaler().fit(X_train)
    X_train_sc, X_val_sc = scaler.transform(X_train), scaler.transform(X_val)

    resultats, modeles_entraines = [], {}
    for nom, modele in construire_modeles().items():
        if nom in MODELES_A_SCALER:
            modele.fit(X_train_sc, y_train)
            y_pred = modele.predict(X_val_sc)
        else:
            modele.fit(X_train, y_train)
            y_pred = modele.predict(X_val)

        rmse = float(np.sqrt(mean_squared_error(y_val, y_pred)))
        mae = float(mean_absolute_error(y_val, y_pred))
        r2 = float(r2_score(y_val, y_pred))
        correlation = float(np.corrcoef(y_val, y_pred)[0, 1]) if len(y_val) > 1 else np.nan
        dir_acc = float((np.where(y_val >= 0, 1, -1) == np.where(y_pred >= 0, 1, -1)).mean())

        resultats.append({"Modele": nom, "RMSE": rmse, "MAE": mae, "R2": r2,
                           "Correlation": correlation, "Directional_Accuracy": dir_acc})
        modeles_entraines[nom] = modele

    df_resultats = pd.DataFrame(resultats).set_index("Modele")
    best_name = df_resultats["Correlation"].idxmax()
    best_needs_scaling = best_name in MODELES_A_SCALER

    # Feature importance (modeles entraines sur train uniquement)
    lignes_fimp = []
    for nom in MODELES_ARBRES:
        raw_imp = np.asarray(modeles_entraines[nom].feature_importances_, dtype=float)
        imp_pct = 100 * raw_imp / raw_imp.sum() if raw_imp.sum() > 0 else np.zeros_like(raw_imp)
        for feat, val in zip(all_features, imp_pct):
            lignes_fimp.append({"Modele": nom, "Feature": feat, "Importance_%": val})
    df_fimp = pd.DataFrame(lignes_fimp)
    df_fimp_moyenne = df_fimp.groupby("Feature")["Importance_%"].mean().sort_values(ascending=False)

    return best_name, best_needs_scaling, df_resultats, df_fimp_moyenne


def allocation_equiponderee(n):
    """Allocation lineaire EQUIPONDEREE : chaque action du Top N recoit 1/N du capital."""
    return np.full(n, 1.0 / n)


# ==============================================================================
# BACKTEST WALK-FORWARD (annee entierement realisee : validation ou test)
# ==============================================================================
@st.cache_data(show_spinner="Backtest walk-forward en cours (reentrainement mensuel)...")
def backtest_walk_forward(df_ml, all_features, annee_cible, nom_modele, needs_scaling, top_n, capital_initial):
    mois_investissement = pd.date_range(f"{annee_cible}-01-01", f"{annee_cible}-12-01", freq="MS")
    mois_decision = mois_investissement - pd.DateOffset(months=1)

    lignes_allocation, lignes_perf, lignes_pred = [], [], []
    valeur_portefeuille = capital_initial

    for date_decision, date_investissement in zip(mois_decision, mois_investissement):
        decision_rows = df_ml[df_ml["Date"] == date_decision].copy()
        if decision_rows.empty:
            continue

        limite_train = date_decision - pd.DateOffset(months=1)
        pool = df_ml[df_ml["Date"] <= limite_train]
        if len(pool) < 30:
            continue

        X_pool, y_pool = pool[all_features], pool[TARGET_COL]
        X_decision = decision_rows[all_features]

        modele_mois = clone(construire_modeles()[nom_modele])
        if needs_scaling:
            scaler_mois = StandardScaler().fit(X_pool)
            modele_mois.fit(scaler_mois.transform(X_pool), y_pool)
            y_pred = modele_mois.predict(scaler_mois.transform(X_decision))
        else:
            modele_mois.fit(X_pool, y_pool)
            y_pred = modele_mois.predict(X_decision)

        decision_rows = decision_rows.assign(Rendement_predit=y_pred).sort_values("Rendement_predit", ascending=False)

        for _, r in decision_rows.iterrows():
            lignes_pred.append({"Ticker": r["Ticker"], "Date_decision": date_decision,
                                 "Date_investissement": date_investissement,
                                 "Rendement_M1_reel": r[TARGET_COL], "Rendement_M1_predit": r["Rendement_predit"]})

        n_disponibles = len(decision_rows)
        if n_disponibles < top_n:
            continue

        topn = decision_rows.head(top_n).reset_index(drop=True)
        poids = allocation_equiponderee(top_n)

        capital_debut = valeur_portefeuille
        rendement_reel_pf = float(np.dot(poids, topn[TARGET_COL].values))
        rendement_predit_pf = float(np.dot(poids, topn["Rendement_predit"].values))
        # Benchmark proxy : rendement moyen EQUIPONDERE de TOUTES les actions disponibles ce
        # mois-la (pas seulement le Top N). PAS le vrai indice MASI/MASI20 (absent des donnees).
        rendement_benchmark = float(decision_rows[TARGET_COL].mean())

        for i in range(top_n):
            lignes_allocation.append({
                "Date_decision": date_decision, "Date_investissement": date_investissement,
                "Ticker": topn.loc[i, "Ticker"], "Rang": i + 1,
                "Rendement_predit_M1": topn.loc[i, "Rendement_predit"], "Allocation": poids[i],
                "Capital_alloue_MAD": capital_debut * poids[i], "Rendement_reel_M1": topn.loc[i, TARGET_COL],
                "Contribution_portefeuille": poids[i] * topn.loc[i, TARGET_COL], "Modele_utilise": nom_modele,
                "Nb_actions_disponibles": n_disponibles,
            })

        valeur_portefeuille = capital_debut * (1 + rendement_reel_pf)
        lignes_perf.append({
            "Date_decision": date_decision, "Date_investissement": date_investissement,
            "Capital_debut_MAD": capital_debut, "Rendement_predit_portefeuille": rendement_predit_pf,
            "Rendement_reel_portefeuille": rendement_reel_pf,
            "Rendement_benchmark_proxy": rendement_benchmark,
            "Surperformance_vs_benchmark": rendement_reel_pf - rendement_benchmark,
            "Gain_perte_MAD": capital_debut * rendement_reel_pf, "Valeur_portefeuille_MAD": valeur_portefeuille,
        })

    df_allocation = pd.DataFrame(lignes_allocation)
    df_perf = pd.DataFrame(lignes_perf)
    df_pred = pd.DataFrame(lignes_pred)
    if len(df_perf):
        df_perf["Rendement_cumule"] = df_perf["Valeur_portefeuille_MAD"] / capital_initial - 1

    return df_allocation, df_perf, df_pred


# ==============================================================================
# BACKTEST + PREVISION SUR L'ANNEE EN COURS (mois realises + mois a venir)
# ==============================================================================
@st.cache_data(show_spinner="Analyse de l'annee en cours (backtest + prevision)...")
def backtest_et_prevision_live(df_features, df_ml, all_features, annee_cible, nom_modele, needs_scaling, top_n, capital_initial):
    mois_investissement = pd.date_range(f"{annee_cible}-01-01", f"{annee_cible}-12-01", freq="MS")
    mois_decision = mois_investissement - pd.DateOffset(months=1)

    lignes_allocation, lignes_perf, lignes_pred, lignes_prevision = [], [], [], []
    valeur_portefeuille = capital_initial

    for date_decision, date_investissement in zip(mois_decision, mois_investissement):
        decision_rows = df_features[df_features["Date"] == date_decision].copy()
        if decision_rows.empty:
            continue

        limite_train = date_decision - pd.DateOffset(months=1)
        pool = df_ml[df_ml["Date"] <= limite_train]
        if len(pool) < 30:
            continue

        X_pool, y_pool = pool[all_features], pool[TARGET_COL]
        X_decision = decision_rows[all_features]

        modele_mois = clone(construire_modeles()[nom_modele])
        if needs_scaling:
            scaler_mois = StandardScaler().fit(X_pool)
            modele_mois.fit(scaler_mois.transform(X_pool), y_pool)
            y_pred = modele_mois.predict(scaler_mois.transform(X_decision))
        else:
            modele_mois.fit(X_pool, y_pool)
            y_pred = modele_mois.predict(X_decision)

        decision_rows = decision_rows.assign(Rendement_predit=y_pred).sort_values("Rendement_predit", ascending=False)
        n_disponibles = len(decision_rows)
        rendement_deja_connu = decision_rows[TARGET_COL].notna().all()

        if rendement_deja_connu:
            for _, r in decision_rows.iterrows():
                lignes_pred.append({"Ticker": r["Ticker"], "Date_decision": date_decision,
                                     "Date_investissement": date_investissement,
                                     "Rendement_M1_reel": r[TARGET_COL], "Rendement_M1_predit": r["Rendement_predit"]})
            if n_disponibles < top_n:
                continue
            topn = decision_rows.head(top_n).reset_index(drop=True)
            poids = allocation_equiponderee(top_n)

            capital_debut = valeur_portefeuille
            rendement_reel_pf = float(np.dot(poids, topn[TARGET_COL].values))
            rendement_predit_pf = float(np.dot(poids, topn["Rendement_predit"].values))
            rendement_benchmark = float(decision_rows[TARGET_COL].mean())

            for i in range(top_n):
                lignes_allocation.append({
                    "Date_decision": date_decision, "Date_investissement": date_investissement,
                    "Ticker": topn.loc[i, "Ticker"], "Rang": i + 1,
                    "Rendement_predit_M1": topn.loc[i, "Rendement_predit"], "Allocation": poids[i],
                    "Capital_alloue_MAD": capital_debut * poids[i], "Rendement_reel_M1": topn.loc[i, TARGET_COL],
                    "Contribution_portefeuille": poids[i] * topn.loc[i, TARGET_COL], "Modele_utilise": nom_modele,
                    "Nb_actions_disponibles": n_disponibles,
                })

            valeur_portefeuille = capital_debut * (1 + rendement_reel_pf)
            lignes_perf.append({
                "Date_decision": date_decision, "Date_investissement": date_investissement,
                "Capital_debut_MAD": capital_debut, "Rendement_predit_portefeuille": rendement_predit_pf,
                "Rendement_reel_portefeuille": rendement_reel_pf,
                "Rendement_benchmark_proxy": rendement_benchmark,
                "Surperformance_vs_benchmark": rendement_reel_pf - rendement_benchmark,
                "Gain_perte_MAD": capital_debut * rendement_reel_pf, "Valeur_portefeuille_MAD": valeur_portefeuille,
            })
        else:
            if n_disponibles < top_n:
                continue
            topn = decision_rows.head(top_n).reset_index(drop=True)
            poids = allocation_equiponderee(top_n)
            for i in range(top_n):
                lignes_prevision.append({
                    "Date_decision": date_decision, "Date_investissement": date_investissement,
                    "Ticker": topn.loc[i, "Ticker"], "Rang": i + 1,
                    "Rendement_predit_M1": topn.loc[i, "Rendement_predit"], "Allocation_equiponderee": poids[i],
                    "Modele_utilise": nom_modele,
                })

    df_allocation = pd.DataFrame(lignes_allocation)
    df_perf = pd.DataFrame(lignes_perf)
    df_pred = pd.DataFrame(lignes_pred)
    df_previsions = pd.DataFrame(lignes_prevision)
    if len(df_perf):
        df_perf["Rendement_cumule"] = df_perf["Valeur_portefeuille_MAD"] / capital_initial - 1

    return df_allocation, df_perf, df_pred, df_previsions


def construire_classement(df_pred):
    if df_pred.empty:
        return pd.DataFrame()
    df_pred = df_pred.copy()
    df_pred["Direction_reelle"] = np.where(df_pred["Rendement_M1_reel"] >= 0, "Hausse", "Baisse")
    df_pred["Direction_predite"] = np.where(df_pred["Rendement_M1_predit"] >= 0, "Hausse", "Baisse")
    classement = (
        df_pred.groupby("Ticker")
        .apply(lambda g: pd.Series({
            "Rendement_predit_moyen": g["Rendement_M1_predit"].mean(),
            "Rendement_reel_moyen": g["Rendement_M1_reel"].mean(),
            "Precision_directionnelle": (g["Direction_predite"] == g["Direction_reelle"]).mean(),
            "Nb_observations": len(g),
        }))
        .reset_index().sort_values("Rendement_predit_moyen", ascending=False).reset_index(drop=True)
    )
    classement.index = classement.index + 1
    classement.index.name = "Rang"
    return classement


def calculer_metriques_portefeuille(df_perf, capital_initial, taux_sans_risque):
    if df_perf.empty:
        return {}
    rm = df_perf["Rendement_reel_portefeuille"].values
    n_mois = len(rm)
    valeur_finale = df_perf["Valeur_portefeuille_MAD"].iloc[-1]

    courbe = np.concatenate([[capital_initial], df_perf["Valeur_portefeuille_MAD"].values])
    sommet = np.maximum.accumulate(courbe)
    max_dd = float((courbe / sommet - 1).min())

    ecart_type = np.std(rm, ddof=1) if n_mois > 1 else np.nan
    sharpe = ((np.mean(rm) - taux_sans_risque / 12) / ecart_type * np.sqrt(12)) if ecart_type and ecart_type > 0 else np.nan

    return {
        "Capital_final_MAD": valeur_finale,
        "Rendement_cumule": valeur_finale / capital_initial - 1,
        "Volatilite_annualisee": float(ecart_type * np.sqrt(12)) if ecart_type == ecart_type else np.nan,
        "Sharpe_Ratio": sharpe,
        "Max_Drawdown": max_dd,
        "Mois_positifs": int((rm > 0).sum()),
        "Mois_negatifs": int((rm < 0).sum()),
    }


# ==============================================================================
# INTERFACE STREAMLIT
# ==============================================================================
st.title("Recommandations de portefeuille ML - Actions MASI20")
st.caption("Reprend exactement la logique du pipeline (donnees -> modeles -> backtest walk-forward -> "
           "benchmark proxy -> prevision) sans passer par Jupyter.")

if not PACKAGES_OK:
    st.error(f"Package manquant : {IMPORT_ERROR}. Ajoutez-le a requirements.txt.")
    st.stop()

with st.sidebar:
    st.header("1. Donnees")
    fichier = st.file_uploader("Fichier Excel (feuilles 'Technique' et 'Fondamental')", type=["xlsx"])

    st.header("2. Parametres")
    top_n = st.slider("Nombre d'actions selectionnees (Top N)", 3, 15, TOP_N_DEFAULT)
    capital_initial = st.number_input("Capital initial (MAD)", value=1_000_000, step=100_000)
    taux_sans_risque = st.number_input("Taux sans risque annuel (%)", value=0.0, step=0.5) / 100

    st.header("3. Periodes")
    annee_train_max = st.number_input("Training jusqu'a (inclus)", value=2023, step=1)
    annee_validation = st.number_input("Annee de validation (selection du modele + 1er backtest)", value=2024, step=1)
    annee_test = st.number_input("Annee de test final (backtest hors-echantillon)", value=2025, step=1)
    annee_live = st.number_input("Annee en cours (backtest mois realises + prevision mois a venir)", value=2026, step=1)
    st.caption("Le modele est choisi UNIQUEMENT sur la validation. Le test et l'annee en cours ne servent "
               "jamais a choisir le modele, seulement a l'evaluer.")

    lancer = st.button("Lancer l'analyse", type="primary", disabled=fichier is None)

if fichier is None:
    st.info("Chargez un fichier Excel dans la barre laterale pour commencer.")
    st.stop()

if not lancer and "resultats_prets" not in st.session_state:
    st.info("Reglez les parametres puis cliquez sur 'Lancer l'analyse'.")
    st.stop()

if lancer:
    st.session_state["resultats_prets"] = True

# ---- Pipeline ----
df_features, df_ml, all_features, tickers_source, tickers_fond, tickers_exclus = preparer_donnees(fichier.getvalue())
best_name, best_needs_scaling, df_resultats, df_fimp_moyenne = selectionner_modele(
    df_ml, all_features, annee_train_max, annee_validation
)

df_alloc_val, df_perf_val, df_pred_val = backtest_walk_forward(
    df_ml, all_features, annee_validation, best_name, best_needs_scaling, top_n, capital_initial
)
df_alloc_test, df_perf_test, df_pred_test = backtest_walk_forward(
    df_ml, all_features, annee_test, best_name, best_needs_scaling, top_n, capital_initial
)
df_alloc_live, df_perf_live, df_pred_live, df_previsions_live = backtest_et_prevision_live(
    df_features, df_ml, all_features, annee_live, best_name, best_needs_scaling, top_n, capital_initial
)

df_classement_val = construire_classement(df_pred_val)
df_classement_test = construire_classement(df_pred_test)
df_classement_live = construire_classement(df_pred_live)

metriques_val = calculer_metriques_portefeuille(df_perf_val, capital_initial, taux_sans_risque)
metriques_test = calculer_metriques_portefeuille(df_perf_test, capital_initial, taux_sans_risque)
metriques_live = calculer_metriques_portefeuille(df_perf_live, capital_initial, taux_sans_risque)

mois_realises = sorted(df_pred_live["Date_investissement"].unique()) if len(df_pred_live) else []
mois_a_venir = sorted(df_previsions_live["Date_investissement"].unique()) if len(df_previsions_live) else []

extremes = df_ml[(df_ml["Annee"].isin([annee_validation, annee_test, annee_live])) &
                  (df_ml[TARGET_COL].abs() > 0.50)][["Ticker", "Date", "Price", TARGET_COL]]
tickers_test = sorted(df_ml.loc[df_ml["Annee"] == annee_test, "Ticker"].unique())

# ---- Affichage ----
onglet_val, onglet_test, onglet_live, onglet_modeles, onglet_diag = st.tabs(
    [f"Validation {annee_validation}", f"Test {annee_test}", f"Annee en cours {annee_live}",
     "Modeles & features", "Diagnostic donnees"]
)


def afficher_backtest(df_classement, df_alloc, df_perf, metriques, annee_cible, capital_initial):
    if metriques:
        c1, c2, c3, c4 = st.columns(4)
        c1.metric("Capital final", f"{metriques['Capital_final_MAD']:,.0f} MAD")
        c2.metric("Rendement cumule", f"{metriques['Rendement_cumule']:.1%}")
        c3.metric("Sharpe Ratio", f"{metriques['Sharpe_Ratio']:.2f}")
        c4.metric("Max Drawdown", f"{metriques['Max_Drawdown']:.1%}")
        if len(df_perf) < 12:
            st.info(f"Seulement {len(df_perf)}/12 mois disponibles pour {annee_cible}.")

        courbe = pd.concat([
            pd.DataFrame({"Date": [df_perf['Date_decision'].iloc[0]], "Valeur": [capital_initial]}),
            df_perf[["Date_investissement", "Valeur_portefeuille_MAD"]].rename(
                columns={"Date_investissement": "Date", "Valeur_portefeuille_MAD": "Valeur"})
        ])
        st.line_chart(courbe.set_index("Date"))

        st.caption("Portefeuille vs benchmark proxy (rendement moyen equipondere de tout l'univers disponible - "
                    "PAS le vrai indice MASI/MASI20, absent des donnees)")
        st.line_chart(df_perf.set_index("Date_investissement")[["Rendement_reel_portefeuille", "Rendement_benchmark_proxy"]])

        st.subheader("Classement")
        st.dataframe(df_classement, use_container_width=True)
        st.subheader("Allocation mensuelle (Top N equipondere)")
        st.dataframe(df_alloc, use_container_width=True, hide_index=True)
        st.subheader("Performance du portefeuille (sans le rendement predit, sur demande)")
        st.dataframe(df_perf.drop(columns=["Rendement_predit_portefeuille"]), use_container_width=True, hide_index=True)
    else:
        st.warning("Backtest impossible avec ces parametres (pas assez de donnees pour cette annee).")


with onglet_val:
    st.caption("Sert a choisir le meilleur modele ET donne une 1ere estimation de performance - "
               "a considerer comme du reporting complementaire, pas la preuve finale.")
    afficher_backtest(df_classement_val, df_alloc_val, df_perf_val, metriques_val, annee_validation, capital_initial)

with onglet_test:
    st.caption(f"Backtest hors-echantillon : le modele {best_name} n'a jamais ete choisi ni ajuste sur {annee_test}.")
    afficher_backtest(df_classement_test, df_alloc_test, df_perf_test, metriques_test, annee_test, capital_initial)

with onglet_live:
    st.caption(f"Mois realises ({', '.join(pd.Timestamp(m).strftime('%b-%Y') for m in mois_realises) or 'aucun'}) : "
               "backtest complet (prediction vs reel). Mois a venir "
               f"({', '.join(pd.Timestamp(m).strftime('%b-%Y') for m in mois_a_venir) or 'aucun'}) : "
               "prevision seule, aucun reel a comparer.")
    st.markdown("#### Mois deja realises")
    afficher_backtest(df_classement_live, df_alloc_live, df_perf_live, metriques_live, annee_live, capital_initial)

    st.markdown("#### Mois a venir (prevision seule)")
    if len(df_previsions_live):
        st.dataframe(df_previsions_live, use_container_width=True, hide_index=True)
        st.warning("Ce sont des PREVISIONS prospectives : aucun rendement reel n'existe encore pour les comparer. "
                   "Ce n'est pas un conseil en investissement.")
    else:
        st.info("Aucun mois a venir disponible - soit toute l'annee est deja realisee, soit les features "
                "(technique + fondamental) ne sont pas encore completes pour les mois suivants.")

with onglet_modeles:
    st.subheader(f"Comparaison des modeles (validation {annee_validation} uniquement)")
    st.dataframe(df_resultats.style.highlight_max(subset=["Correlation"], color="lightgreen"), use_container_width=True)
    st.subheader("Importance moyenne des variables")
    st.caption("Calculee automatiquement par les modeles a base d'arbres - non fixee manuellement dans le code.")
    st.bar_chart(df_fimp_moyenne)

with onglet_diag:
    st.subheader("Univers de tickers")
    st.write(f"{len(tickers_source)} tickers dans la source -> {len(tickers_test)} utilisables en {annee_test}.")
    st.write(f"Tickers exclus : {', '.join(tickers_exclus) if tickers_exclus else 'aucun'}")
    st.subheader(f"Rendements mensuels extremes (> 50%, {annee_validation}-{annee_live}, a verifier manuellement)")
    st.dataframe(extremes, use_container_width=True, hide_index=True)

# ==============================================================================
# EXPORT EXCEL - MEME STRUCTURE QUE LE PIPELINE (colonnes ajustees, %, MAD)
# ==============================================================================
HEADER_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
HEADER_FONT = Font(bold=True, color="FFFFFF")

def ecrire_feuille_formatee(wb, nom_feuille, df, pct_cols=(), mad_cols=()):
    ws = wb.create_sheet(nom_feuille[:31])
    if df is None or df.empty:
        ws["A1"] = "(aucune donnee disponible)"
        return ws
    for j, col in enumerate(df.columns, start=1):
        c = ws.cell(row=1, column=j, value=str(col))
        c.font = HEADER_FONT
        c.fill = HEADER_FILL
        c.alignment = Alignment(horizontal="center", vertical="center")
    for i, (_, row) in enumerate(df.iterrows(), start=2):
        for j, col in enumerate(df.columns, start=1):
            val = row[col]
            if isinstance(val, pd.Timestamp):
                val = val.to_pydatetime()
            cell = ws.cell(row=i, column=j, value=val)
            if col in pct_cols:
                cell.number_format = "0.00%"
            elif col in mad_cols:
                cell.number_format = '#,##0 "MAD"'
            elif "Date" in str(col):
                cell.number_format = "dd/mm/yyyy"
    for j, col in enumerate(df.columns, start=1):
        contenu_max = df[col].astype(str).map(len).max() if len(df) else 0
        largeur = max(len(str(col)), contenu_max) + 3
        ws.column_dimensions[get_column_letter(j)].width = min(largeur, 45)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(df.columns))}{len(df) + 1}"
    ws.row_dimensions[1].height = 20
    return ws


PERF_PCT_COLS = ["Rendement_reel_portefeuille", "Rendement_benchmark_proxy", "Surperformance_vs_benchmark", "Rendement_cumule"]
PERF_MAD_COLS = ["Capital_debut_MAD", "Gain_perte_MAD", "Valeur_portefeuille_MAD"]
ALLOC_PCT_COLS = ["Rendement_predit_M1", "Allocation", "Rendement_reel_M1", "Contribution_portefeuille"]

st.divider()
wb_export = Workbook()
wb_export.remove(wb_export.active)

ecrire_feuille_formatee(wb_export, "Comparaison_Modeles", df_resultats.reset_index(), pct_cols=["Directional_Accuracy"])
ecrire_feuille_formatee(wb_export, f"Classement_Valid_{annee_validation}", df_classement_val.reset_index(),
                         pct_cols=["Rendement_predit_moyen", "Rendement_reel_moyen", "Precision_directionnelle"])
ecrire_feuille_formatee(wb_export, f"Alloc_Valid_{annee_validation}", df_alloc_val, pct_cols=ALLOC_PCT_COLS, mad_cols=["Capital_alloue_MAD"])
ecrire_feuille_formatee(wb_export, f"Perf_Valid_{annee_validation}",
                         df_perf_val.drop(columns=["Rendement_predit_portefeuille"], errors="ignore"),
                         pct_cols=PERF_PCT_COLS, mad_cols=PERF_MAD_COLS)
ecrire_feuille_formatee(wb_export, f"Classement_Test_{annee_test}", df_classement_test.reset_index(),
                         pct_cols=["Rendement_predit_moyen", "Rendement_reel_moyen", "Precision_directionnelle"])
ecrire_feuille_formatee(wb_export, f"Alloc_Test_{annee_test}", df_alloc_test, pct_cols=ALLOC_PCT_COLS, mad_cols=["Capital_alloue_MAD"])
ecrire_feuille_formatee(wb_export, f"Perf_Test_{annee_test}",
                         df_perf_test.drop(columns=["Rendement_predit_portefeuille"], errors="ignore"),
                         pct_cols=PERF_PCT_COLS, mad_cols=PERF_MAD_COLS)
ecrire_feuille_formatee(wb_export, f"Classement_{annee_live}", df_classement_live.reset_index(),
                         pct_cols=["Rendement_predit_moyen", "Rendement_reel_moyen", "Precision_directionnelle"])
ecrire_feuille_formatee(wb_export, f"Alloc_{annee_live}", df_alloc_live, pct_cols=ALLOC_PCT_COLS, mad_cols=["Capital_alloue_MAD"])
ecrire_feuille_formatee(wb_export, f"Perf_{annee_live}",
                         df_perf_live.drop(columns=["Rendement_predit_portefeuille"], errors="ignore"),
                         pct_cols=PERF_PCT_COLS, mad_cols=PERF_MAD_COLS)
ecrire_feuille_formatee(wb_export, f"Prevision_a_venir_{annee_live}", df_previsions_live,
                         pct_cols=["Rendement_predit_M1", "Allocation_equiponderee"])

buffer = io.BytesIO()
wb_export.save(buffer)
st.download_button("Telecharger le rapport Excel", data=buffer.getvalue(),
                    file_name="Resultats_ML_Portefeuille.xlsx",
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
